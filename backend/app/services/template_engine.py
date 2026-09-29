import io
import re
from typing import Dict, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

class TemplateEngine:
    @staticmethod
    def render_text(template_content: str, variables: Dict[str, Any]) -> str:
        rendered = template_content
        for key, value in variables.items():
            pattern = re.compile(r'\{\{\s*' + re.escape(key) + r'\s*\}\}', re.IGNORECASE)
            rendered = pattern.sub(str(value if value is not None else ''), rendered)
        return rendered

    @staticmethod
    def generate_docx(
        title: str,
        content: str,
        company_name: str = "BizBook Enterprise",
        letterhead_header: Optional[str] = None,
        letterhead_footer: Optional[str] = None,
        signature_name: Optional[str] = None,
        signature_designation: Optional[str] = None,
        seal_name: Optional[str] = None,
    ) -> bytes:
        doc = docx.Document()

        # Header / Letterhead
        header = doc.sections[0].header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(letterhead_header or f"OFFICIAL CORPORATE DOCUMENT • {company_name.upper()}")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(128, 128, 128)

        # Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_title = p_title.add_run(title)
        run_title.bold = True
        run_title.font.size = Pt(18)
        run_title.font.color.rgb = RGBColor(30, 58, 138)

        # Body paragraphs
        for para_text in content.split("\n\n"):
            if para_text.strip():
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run(para_text.strip())
                run.font.size = Pt(11)

        # Signature & Seal Block
        if signature_name or seal_name:
            doc.add_paragraph().paragraph_format.space_before = Pt(24)
            table = doc.add_table(rows=2, cols=2)
            table.autofit = True
            
            # Left: Signature
            cell_sig = table.cell(0, 0)
            cell_sig.paragraphs[0].add_run("[ Authorized Digital Signature Applied ]\n").font.size = Pt(9)
            if signature_name:
                r_name = cell_sig.paragraphs[0].add_run(f"\n{signature_name}\n")
                r_name.bold = True
                r_name.font.size = Pt(10)
            if signature_designation:
                cell_sig.paragraphs[0].add_run(f"{signature_designation}\n").font.size = Pt(9)

            # Right: Seal
            cell_seal = table.cell(0, 1)
            if seal_name:
                cell_seal.paragraphs[0].add_run(f"[ Verified Corporate Seal: {seal_name} ]").font.size = Pt(9)

        # Footer
        footer = doc.sections[0].footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run(letterhead_footer or f"BizBook Verified Institutional Memory • Generated for {company_name}")
        frun.font.size = Pt(8)
        frun.font.color.rgb = RGBColor(128, 128, 128)

        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()
